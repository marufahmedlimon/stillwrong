"""StillWrong evidence extraction, integrity validation and static reporting."""
import argparse
import csv
import hashlib
import html
import json
import math
import os
import re
import stat
import sys
import tempfile
import zipfile
from decimal import Decimal
from fractions import Fraction
from string import Template
from pathlib import Path

VERSION = '0.1.1'
SUPPORTED = {'.txt', '.md', '.csv', '.pdf', '.docx', '.pptx', '.xlsx'}
MAX_FILE = 25 * 1024 * 1024
MAX_TEXT = 1_000_000
MAX_RECORDS = 50_000
MAX_CELLS = 100_000
MAX_ENTRIES = 500
EXCLUDED = {'.git', '.venv', 'venv', '__pycache__', 'node_modules'}

class LimitError(ValueError):
    """A declared work limit was reached; coverage is incomplete."""

def is_link(path):
    info = path.lstat()
    return path.is_symlink() or bool(getattr(info, 'st_file_attributes', 0) & 0x400)

def dump(path, data):
    Path(path).write_text(json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False), encoding='utf-8')

def load(path):
    return json.loads(Path(path).read_text(encoding='utf-8-sig'))

def digest(path):
    result = hashlib.sha256()
    with Path(path).open('rb') as stream:
        for block in iter(lambda: stream.read(65536), b''):
            result.update(block)
    return result.hexdigest()

def read_file(path, text_limit=MAX_TEXT):
    records, notices = [], []
    text_size = 0
    def add(locator, text):
        nonlocal text_size
        if text is None:
            return
        text = str(text)
        if text.strip():
            text_size += len(text)
            if text_size > text_limit or len(records) >= MAX_RECORDS:
                raise LimitError('Text or record limit reached; narrow the selected input')
            records.append({'locator': locator, 'text': text})
    ext = path.suffix.lower()
    if ext in {'.docx', '.pptx', '.xlsx'}:
        with zipfile.ZipFile(path) as z:
            if sum(i.file_size for i in z.infolist()) > 100 * 1024 * 1024:
                raise ValueError('Expanded archive exceeds 100 MB')
            names = z.namelist()
            if any('/media/' in n or '/embeddings/' in n or '/charts/' in n or '/diagrams/' in n for n in names):
                notices.append('Images, charts, diagrams or embedded objects are not reviewed')
            if ext == '.docx':
                raw = b''.join(z.read(n) for n in names if n.startswith('word/') and n.endswith('.xml'))
                if any(tag in raw for tag in [b'<w:ins', b'<w:del', b'<w:txbxContent', b'<w:footnoteReference']):
                    notices.append('Tracked revisions, textboxes or footnotes not reviewed')
    if ext in {'.txt', '.md'}:
        with path.open(encoding='utf-8-sig') as stream:
            for n, line in enumerate(stream, 1):
                add(f'line {n}', line.rstrip('\r\n'))
    elif ext == '.csv':
        with path.open(encoding='utf-8-sig', newline='') as f:
            for row, cells in enumerate(csv.reader(f), 1):
                for col, value in enumerate(cells, 1):
                    add(f'record {row}/column {col}', value)
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(path)
        if reader.is_encrypted:
            raise ValueError('Encrypted PDF; supply an unlocked copy')
        if len(reader.pages) > 500:
            raise LimitError('PDF exceeds 500-page limit')
        for n, page in enumerate(reader.pages, 1):
            resources = page.get('/Resources', {})
            if hasattr(resources, 'get_object'):
                resources = resources.get_object()
            if resources.get('/XObject'):
                notices.append(f'Page {n}: images not reviewed; OCR unavailable')
            text = page.extract_text() or ''
            if not text.strip():
                notices.append(f'Page {n}: no extractable text')
            add(f'page {n}', text)
        notices.append('PDF text order/table structure not visually verified')
    elif ext == '.docx':
        from docx import Document
        doc = Document(path)
        def walk(container, prefix, depth=0):
            if depth > 30:
                raise LimitError('Nested table depth exceeds 30')
            for n, p in enumerate(container.paragraphs, 1):
                add(f'{prefix}/paragraph {n}', p.text)
            for n, table in enumerate(container.tables, 1):
                for r, row in enumerate(table.rows, 1):
                    for c, cell in enumerate(row.cells, 1):
                        walk(cell, f'{prefix}/table {n}/row {r}/cell {c}', depth+1)
        walk(doc, 'body')
        for n, section in enumerate(doc.sections, 1):
            for name in ['header', 'footer', 'first_page_header', 'first_page_footer', 'even_page_header', 'even_page_footer']:
                walk(getattr(section, name), f'section {n}/{name}')
        notices.append('Word fields, comments, drawings and revision history not semantically reviewed')
    elif ext == '.pptx':
        from pptx import Presentation
        from pptx.enum.shapes import MSO_SHAPE_TYPE
        def shapes(items, prefix, depth=0):
            if depth > 30:
                raise LimitError('Grouped shape depth exceeds 30')
            for n, shape in enumerate(items, 1):
                loc = f'{prefix}/shape {n}'
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    shapes(shape.shapes, loc, depth+1)
                if shape.has_text_frame:
                    add(loc, shape.text)
                if shape.has_table:
                    for r, row in enumerate(shape.table.rows, 1):
                        for c, cell in enumerate(row.cells, 1):
                            add(f'{loc}/row {r}/cell {c}', cell.text)
        for n, slide in enumerate(Presentation(path).slides, 1):
            shapes(slide.shapes, f'slide {n}')
        notices.append('Slide masters, layouts and speaker notes not reviewed')
    elif ext == '.xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, data_only=False, read_only=True, keep_links=False)
        cache = load_workbook(path, data_only=True, read_only=True, keep_links=False)
        try:
            for sheet in book:
                if (sheet.max_row or 0) * (sheet.max_column or 0) > MAX_CELLS:
                    raise LimitError(f'{sheet.title}: worksheet dimensions exceed {MAX_CELLS} cells')
                cached_rows = cache[sheet.title].iter_rows()
                visited = 0
                for row, cached_row in zip(sheet.iter_rows(), cached_rows):
                    visited += max(1, len(row))
                    if visited > MAX_CELLS:
                        raise LimitError(f'{sheet.title}: actual iteration exceeds {MAX_CELLS} cells/rows')
                    for cell, cached_cell in zip(row, cached_row):
                        if cell.value is not None:
                            loc = f'{sheet.title}!{cell.coordinate}'
                            if cell.data_type == 'f':
                                if isinstance(cell.value, str):
                                    formula = cell.value
                                else:
                                    formula = json.dumps(dict(kind=type(cell.value).__name__, **vars(cell.value)), sort_keys=True)
                                    notices.append(f'{loc}: special formula interpretation unsupported')
                                add(loc, formula)
                                add(loc + '/cached (potentially stale)', cached_cell.value)
                                if cached_cell.value is None:
                                    notices.append(f'{loc}: cached result missing')
                                notices.append(f'{loc}: formula not recalculated')
                            else:
                                add(loc, cell.value)
        finally:
            book.close()
            cache.close()
        notices.append('Workbook charts, comments, drawings and external references not reviewed')
    return records, sorted(set(notices))

def extract(source):
    source = Path(source).absolute()
    if is_link(source):
        raise ValueError('Select a real file or folder, not a link')
    source = source.resolve()
    if not source.exists():
        raise ValueError('Input does not exist')
    root = source if source.is_dir() else source.parent
    files, records = [], []
    candidates = inventory(source)
    total = 0
    for path, special in candidates:
        rel = path.relative_to(root).as_posix()
        item = {'path': rel, 'status': 'unsupported', 'notices': []}
        files.append(item)
        if special:
            item.update(status=special[0], notices=[special[1]])
            continue
        try:
            if is_link(path) or not path.resolve().is_relative_to(root):
                item.update(status='excluded', notices=['External links not followed'])
                continue
            if not stat.S_ISREG(path.stat().st_mode):
                item.update(status='excluded', notices=['Not a regular file'])
                continue
            if path.stat().st_size > MAX_FILE:
                raise LimitError('File exceeds 25 MiB limit')
            item['sha256'] = digest(path)
            if path.suffix.lower() not in SUPPORTED:
                item['notices'] = ['Unsupported format']
                continue
            rows, notices = read_file(path, MAX_TEXT-total)
            size = sum(len(r['text']) for r in rows)
            if digest(path) != item['sha256']:
                raise ValueError('Source changed during extraction; retry with a stable copy')
            total += size
            if not rows:
                notices.append('No extractable nonempty text; content has not been reviewed')
            item.update(status='partial' if notices or not rows else 'extracted', notices=sorted(set(notices)))
            for row in rows:
                records.append(dict(id=f'r{len(records)+1}', file=rel, **row))
        except LimitError as exc:
            item.update(status='partial', notices=[str(exc)])
        except Exception as exc:
            item.update(status='unreadable', notices=[f'{type(exc).__name__}: {exc}'])
    return {'schema_version': 1, 'version': VERSION, 'source': str(source), 'files': files, 'records': records}

def inventory(source):
    """Prune links and caches before descent; cap directories and entries."""
    if is_link(source):
        raise ValueError('Select a real file or folder, not a link')
    if not source.is_dir():
        yield source, None
        return
    count = 0
    def walk(folder, depth=0):
        nonlocal count
        if depth > 30:
            yield folder, ('partial', 'Directory depth limit reached')
            return
        try:
            entries = []
            with os.scandir(folder) as stream:
                for entry in stream:
                    entries.append(Path(entry.path))
                    if len(entries) > MAX_ENTRIES:
                        yield folder, ('partial', 'Directory inventory exceeds 500 entries; narrow selection')
                        return
        except OSError as exc:
            yield folder, ('unreadable', f'Cannot list directory: {exc}')
            return
        for path in sorted(entries):
            count += 1
            if count > MAX_ENTRIES:
                yield folder, ('partial', 'Inventory exceeds 500 entries; narrow selection')
                return
            try:
                if is_link(path):
                    yield path, ('excluded', 'Links and Windows reparse points not followed')
                elif path.is_dir():
                    if path.name in EXCLUDED:
                        yield path, ('excluded', 'Dependency/cache directory pruned')
                    else:
                        yield from walk(path, depth+1)
                else:
                    yield path, None
            except OSError as exc:
                yield path, ('unreadable', f'Cannot inspect entry: {exc}')
    yield from walk(source)

def validate(evidence, findings):
    if not isinstance(evidence, dict) or not isinstance(findings, dict):
        raise ValueError('Evidence and findings must be JSON objects')
    if any(type(d.get('schema_version')) is not int or d['schema_version'] != 1 for d in [evidence, findings]):
        raise ValueError('Unsupported schema version')
    fresh = extract(evidence['source'])
    if fresh != evidence:
        raise ValueError('Evidence or source snapshot changed; extract and review again')
    change = findings.get('change')
    if not isinstance(change, dict):
        raise ValueError('change must be an object')
    if not isinstance(findings.get('findings'), list) or not isinstance(findings.get('exclusions', []), list):
        raise ValueError('findings and exclusions must be arrays')
    for key in ['entity', 'fact', 'old', 'new', 'authority']:
        if not isinstance(change.get(key), str) or not change[key].strip():
            raise ValueError(f'Missing change.{key}')
    index = {r['id']: r for r in evidence['records']}
    seen = set()
    for finding in findings['findings']:
        if not isinstance(finding, dict) or finding.get('record_id') not in index:
            raise ValueError('Unknown finding record_id')
        record = index[finding['record_id']]
        quote = finding.get('quote')
        if not isinstance(quote, str) or not quote.strip() or quote not in record['text']:
            raise ValueError('Quote does not resolve to evidence')
        kind = finding.get('classification')
        if kind not in {'stale', 'derived', 'needs_review'}:
            raise ValueError('Invalid classification')
        identity = (finding['record_id'], quote)
        if identity in seen:
            raise ValueError('Duplicate finding for the same quote')
        seen.add(identity)
        if finding.get('correction') is not None and not isinstance(finding['correction'], str):
            raise ValueError('correction must be text or null')
        for key in ['reason', 'authority']:
            if not isinstance(finding.get(key), str) or not finding[key].strip():
                raise ValueError(f'Missing {key}')
        if kind != 'needs_review' and not isinstance(finding.get('correction'), str):
            raise ValueError('Correction required')
        if kind != 'needs_review' and not finding['correction'].strip():
            raise ValueError('Empty correction')
        if kind == 'derived':
            calc = finding.get('calculation')
            if not isinstance(calc, dict) or not isinstance(calc.get('operands'), list) or not 1 <= len(calc['operands']) <= 50:
                raise ValueError('calculation requires 1–50 decimal-string operands')
            vals = [Fraction(decimal_value(v)) for v in calc['operands']]
            expected = Fraction(decimal_value(calc.get('result')))
            op = calc.get('op')
            if op in {'subtract', 'divide'} and len(vals) != 2:
                raise ValueError('Two operands required')
            if op == 'add': result = sum(vals)
            elif op == 'multiply': result = math.prod(vals)
            elif op == 'subtract': result = vals[0] - vals[1]
            elif op == 'divide':
                if vals[1] == 0: raise ValueError('Division by zero')
                result = vals[0] / vals[1]
            else: raise ValueError('Unsupported calculation')
            if result != expected:
                raise ValueError('Incorrect arithmetic')
            refs = calc.get('relationship_record_ids')
            if not isinstance(refs, list) or not refs or any(not isinstance(i, str) or i not in index for i in refs):
                raise ValueError('Relationship evidence required')
    for excluded in findings.get('exclusions', []):
        if not isinstance(excluded, dict) or excluded.get('record_id') not in index or not isinstance(excluded.get('reason'), str) or not excluded['reason'].strip():
            raise ValueError('Invalid exclusion')
    return index

def decimal_value(value):
    if not isinstance(value, str) or len(value) > 128 or not re.fullmatch(r'[+-]?(?:\d+(?:\.\d*)?|\.\d+)(?:[eE][+-]?\d+)?', value):
        raise ValueError('Calculations require bounded decimal strings, not JSON numbers')
    number = Decimal(value)
    if not number.is_finite() or abs(number.as_tuple().exponent) > 100:
        raise ValueError('Nonfinite or excessive decimal exponent')
    return number

def render(evidence, findings, target, overwrite=False):
    index = validate(evidence, findings)
    target = Path(target).absolute()
    if any(p.exists() and is_link(p) for p in [target, *target.parents]):
        raise ValueError('Report output cannot use linked directories')
    target = target.resolve()
    source = Path(evidence['source']).resolve()
    destinations = [target/name for name in ['report.html', 'report.md', 'report.json']]
    if (source.is_dir() and target.is_relative_to(source)) or source in destinations:
        raise ValueError('Report output would overlap the selected source')
    if not target.is_dir():
        raise ValueError('Report output directory must already exist')
    for path in destinations:
        if path.is_symlink() or (path.exists() and is_link(path)):
            raise ValueError('Report destination cannot be a link')
        if path.exists() and not overwrite:
            raise ValueError('Report exists; use --overwrite to replace generated reports')
    esc = lambda value: html.escape(str(value))
    cards, md = [], ['# StillWrong', '', 'One fact changed. Find what\'s still wrong.', '', '## Change', json.dumps(findings['change'], ensure_ascii=False), '', '## Findings']
    for number, f in enumerate(findings['findings'], 1):
        r = index[f['record_id']]
        location = f"{r['file']} — {r['locator']}"
        cards.append(f'<article data-kind="{esc(f["classification"])}"><span>{esc(f["classification"])}</span><h3>{esc(location)}</h3><blockquote>{esc(f["quote"])}</blockquote><p>{esc(f["reason"])}</p><p><b>Correction draft:</b> {esc(f.get("correction") or "Review needed — no correction verified")}</p><details><summary>Evidence & authority</summary><p>{esc(f["authority"])}</p><pre>{esc(r["text"])}</pre></details></article>')
        md.extend(['', f'### Finding {number} · {f["classification"]}', literal(location), '', 'Original:', literal(f['quote']), 'Reason:', literal(f['reason']), 'Authority:', literal(f['authority']), 'Correction draft:', literal(f.get('correction') or 'Review needed; no correction verified')])
        if f.get('calculation'):
            calc = f['calculation']
            symbol = {'add':'+', 'subtract':'−', 'multiply':'×', 'divide':'÷'}[calc['op']]
            proof = f' {symbol} '.join(calc['operands'])+' = '+calc['result']
            relationships = ''.join('<p><b>Relationship evidence:</b> '+esc(index[rid]['file']+' · '+index[rid]['locator'])+'</p><blockquote>'+esc(index[rid]['text'])+'</blockquote>' for rid in calc['relationship_record_ids'])
            cards[-1] = cards[-1].replace('</article>', '<h4>Calculation</h4><pre>'+esc(proof)+'</pre>'+relationships+'</article>')
            md.extend(['Calculation:', literal(proof)])
    coverage = ''.join(f'<li><b>{esc(f["path"])}</b> · {esc(f["status"])}<p>{esc("; ".join(f["notices"]))}</p></li>' for f in evidence['files'])
    exclusions = ''.join(f'<li><b>{esc(index[x["record_id"]]["file"])} · {esc(index[x["record_id"]]["locator"])}</b><pre>{esc(index[x["record_id"]]["text"])}</pre><p>{esc(x["reason"])}</p></li>' for x in findings.get('exclusions', []))
    md[md.index('## Change')+1] = literal(json.dumps(findings['change'], ensure_ascii=False, indent=2))
    md.extend(['', '## Preserved references'] + [literal(f'{x["record_id"]} · {index[x["record_id"]]["file"]} · {index[x["record_id"]]["locator"]}\nOriginal: {index[x["record_id"]]["text"]}\nReason: {x["reason"]}') for x in findings.get('exclusions', [])])
    md.extend(['', '## Extraction coverage'] + [literal(f'{f["path"]}: {f["status"]}; {"; ".join(f["notices"])}') for f in evidence['files']])
    note = 'Evidence integrity validated; semantic judgments require review. No confirmed findings does not certify unreviewed content. Source files were not edited.'
    md.extend(['', note])
    counts = {kind: sum(f['classification'] == kind for f in findings['findings']) for kind in ['stale', 'derived', 'needs_review']}
    md.insert(md.index('## Findings')+1, f'\nStale: {counts["stale"]} · Derived: {counts["derived"]} · Needs review: {counts["needs_review"]}\n')
    template = Template((Path(__file__).resolve().parent.parent/'assets/report.html').read_text(encoding='utf-8'))
    change = findings['change']
    page = template.substitute(entity=esc(change['entity']), fact=esc(change['fact']), old=esc(change['old']), new=esc(change['new']), authority=esc(change['authority']), stale=counts['stale'], derived=counts['derived'], review=counts['needs_review'], files=len(evidence['files']), cards=''.join(cards) or '<p>No confirmed issues in reviewed content.</p>', coverage=coverage, exclusions=exclusions or '<li>No exclusions recorded.</li>', note=esc(note))
    staged = []
    try:
        bundle = dict(schema_version=1, version=VERSION, change=findings['change'], findings=findings['findings'], exclusions=findings.get('exclusions', []), files=evidence['files'], records=evidence['records'])
        for text in [page, '\n'.join(md), json.dumps(bundle, ensure_ascii=False, indent=2, allow_nan=False)]:
            with tempfile.NamedTemporaryFile(mode='w', encoding='utf-8', dir=target, prefix='.stillwrong-', delete=False) as stream:
                stream.write(text)
                staged.append(Path(stream.name))
        for temp, destination in zip(staged, destinations):
            os.replace(temp, destination)
    finally:
        for temp in staged:
            temp.unlink(missing_ok=True)

def literal(value):
    """A variable-length code fence makes arbitrary Markdown inert."""
    text = str(value)
    fence = '`' * max(3, max((len(m) for m in re.findall(r'`+', text)), default=0)+1)
    return f'\n{fence}text\n{text}\n{fence}\n'

def main():
    parser = argparse.ArgumentParser(description=__doc__)
    sub = parser.add_subparsers(dest='command', required=True)
    p = sub.add_parser('extract'); p.add_argument('input'); p.add_argument('--out', required=True)
    for command in ['validate', 'render']:
        p = sub.add_parser(command); p.add_argument('evidence'); p.add_argument('findings')
        if command == 'render': p.add_argument('--overwrite', action='store_true', help='Replace existing reports, never source files')
    args = parser.parse_args()
    try:
        if args.command == 'extract':
            raw_source = Path(args.input).absolute()
            if is_link(raw_source):
                raise ValueError('Select a real file or folder, not a link')
            source, out = raw_source.resolve(), Path(args.out).resolve()
            if out.is_relative_to(source) or out == source:
                raise ValueError('Output must be outside source')
            data = extract(source)
            out.mkdir(parents=True, exist_ok=False)
            dump(out / 'evidence.json', data)
            print(out / 'evidence.json')
        else:
            evidence, findings = load(args.evidence), load(args.findings)
            if args.command == 'validate':
                validate(evidence, findings); print('Evidence and calculations validated; semantic accuracy not certified.')
            else:
                render(evidence, findings, Path(args.evidence).parent, overwrite=args.overwrite)
                print(Path(args.evidence).parent / 'report.html')
    except Exception as exc:
        parser.exit(1, f'{type(exc).__name__}: {exc}\n')

if __name__ == '__main__':
    main()
