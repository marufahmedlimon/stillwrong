import copy
import importlib.util
import tempfile
import unittest
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
spec = importlib.util.spec_from_file_location('sw', ROOT/'skills/stillwrong/scripts/stillwrong.py')
sw = importlib.util.module_from_spec(spec); spec.loader.exec_module(sw)

class IntegrityTests(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory(dir=ROOT); self.addCleanup(self.tmp.cleanup)
        self.root = Path(self.tmp.name)/'source'; self.root.mkdir()
        (self.root/'note.md').write_text('Orion current budget USD 15000\nAdvance is 20% of approved budget.\nHistorical USD 15000\n<script>alert(1)</script>', encoding='utf-8')
        self.ev = sw.extract(self.root)
        self.findings = dict(schema_version=1,change=dict(entity='Orion',fact='budget',old='15000',new='12000',authority='User'),findings=[dict(record_id='r1',quote='Orion current budget USD 15000',classification='stale',reason='User changed budget',authority='User',correction='Orion current budget USD 12000')],exclusions=[])

    def test_valid_evidence(self): sw.validate(self.ev,self.findings)
    def test_fabricated_quote(self):
        self.findings['findings'][0]['quote']='Fabricated'
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)
    def test_changed_source(self):
        (self.root/'note.md').write_text('changed')
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)
    def test_tampered_evidence(self):
        self.ev['records'][0]['text']='tampered'
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)
    def test_calculation(self):
        f=self.findings['findings'][0]; f['classification']='derived'; f['calculation']=dict(op='multiply',operands=['12000','.2'],result='2400',relationship_record_ids=['r2'])
        sw.validate(self.ev,self.findings)
        f['calculation']['result']='3000'
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)
    def test_empty_correction(self):
        self.findings['findings'][0]['correction']=' '
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)
    def test_report_escape_and_preserve(self):
        before=sw.digest(self.root/'note.md'); out=Path(self.tmp.name)/'out'; out.mkdir()
        self.findings['change']['authority']='<script>alert(1)</script>'
        sw.render(self.ev,self.findings,out)
        self.assertIn('&lt;script&gt;', (out/'report.html').read_text(encoding='utf-8'))
        self.assertEqual(before,sw.digest(self.root/'note.md'))
    def test_unsupported_visible(self):
        (self.root/'image.png').write_bytes(b'not image')
        self.assertIn('unsupported',[f['status'] for f in sw.extract(self.root)['files']])
    def test_utf8_csv(self):
        (self.root/'data.csv').write_text('name,value\n\u09ac\u09be\u0982\u09b2\u09be,\u09e7\u09eb\u09e6\u09e6\u09e6\n',encoding='utf-8')
        self.assertTrue(any(r['text']=='\u09e7\u09eb\u09e6\u09e6\u09e6' for r in sw.extract(self.root)['records']))

    def test_docx_nested_table_and_footer(self):
        from docx import Document
        doc=Document(); cell=doc.add_table(rows=1,cols=1).cell(0,0)
        cell.add_table(rows=1,cols=1).cell(0,0).text='Nested budget 15000'
        doc.sections[0].footer.paragraphs[0].text='Footer 15000'
        doc.save(self.root/'test.docx')
        rows=sw.extract(self.root)['records']
        self.assertTrue(any('table 1/row 1/cell 1/table' in r['locator'] and 'Nested' in r['text'] for r in rows))
        self.assertTrue(any('footer' in r['locator'] and 'Footer' in r['text'] for r in rows))

    def test_grouped_slides(self):
        from pptx import Presentation
        from pptx.util import Inches
        deck=Presentation(); slide=deck.slides.add_slide(deck.slide_layouts[6]); group=slide.shapes.add_group_shape()
        group.shapes.add_textbox(Inches(1),Inches(1),Inches(2), Inches(1)).text='Grouped 15000'
        deck.save(self.root/'test.pptx')
        ev=sw.extract(self.root)
        self.assertTrue(any('shape 1/shape 1' in r['locator'] and r['text']=='Grouped 15000' for r in ev['records']))

    def test_formula_cache_disclosed(self):
        from openpyxl import Workbook
        book=Workbook(); book.active['A1']='=12000*0.2'; book.save(self.root/'test.xlsx')
        ev=sw.extract(self.root)
        self.assertTrue(any(r['text']=='=12000*0.2' for r in ev['records']))
        item=next(f for f in ev['files'] if f['path']=='test.xlsx')
        self.assertEqual(item['status'],'partial')
        self.assertTrue(any('not recalculated' in n for n in item['notices']))

    def test_encrypted_pdf(self):
        from pypdf import PdfWriter
        writer=PdfWriter(); writer.add_blank_page(100,100); writer.encrypt('secret')
        with (self.root/'locked.pdf').open('wb') as out: writer.write(out)
        item=next(f for f in sw.extract(self.root)['files'] if f['path']=='locked.pdf')
        self.assertEqual(item['status'],'unreadable')

    def test_unreadable_docx_visible(self):
        (self.root/'broken.docx').write_bytes(b'broken')
        item=next(f for f in sw.extract(self.root)['files'] if f['path']=='broken.docx')
        self.assertEqual(item['status'],'unreadable')

if __name__ == '__main__': unittest.main()
