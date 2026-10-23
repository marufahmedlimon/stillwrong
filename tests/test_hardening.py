"""Regression cases discovered during the publication audit."""
import copy
import json
import os
import io
import re
import sys
import subprocess
import zipfile
import unittest
from pathlib import Path
from unittest.mock import patch
import test_stillwrong as base
sw = base.sw


class HardeningTests(unittest.TestCase):
    setUp = base.IntegrityTests.setUp

    def test_directory_source_cannot_receive_reports(self):
        with self.assertRaisesRegex(ValueError, 'overlap'):
            sw.render(self.ev, self.findings, self.root)

    def test_single_source_named_report_cannot_be_overwritten(self):
        source = self.root/'report.md'
        source.write_text('Original', encoding='utf-8')
        evidence = sw.extract(source)
        findings = copy.deepcopy(self.findings); findings['findings'] = []
        with self.assertRaisesRegex(ValueError, 'overlap'):
            sw.render(evidence, findings, self.root, overwrite=True)
        self.assertEqual(source.read_text(encoding='utf-8'), 'Original')

    def test_existing_report_needs_explicit_overwrite(self):
        out = Path(self.tmp.name)/'out'; out.mkdir()
        sw.render(self.ev, self.findings, out)
        with self.assertRaisesRegex(ValueError, 'exists'):
            sw.render(self.ev, self.findings, out)
        sw.render(self.ev, self.findings, out, overwrite=True)
        self.assertEqual(json.loads((out/'report.json').read_text(encoding='utf-8'))['schema_version'], 1)

    def test_markdown_is_literal_even_with_embedded_fences(self):
        attack = '```\n![remote](https://example.com/pixel)\n<script>attack</script>\n```'
        self.findings['change']['authority'] = attack
        out = Path(self.tmp.name)/'out'; out.mkdir()
        sw.render(self.ev, self.findings, out)
        text = (out/'report.md').read_text(encoding='utf-8')
        self.assertIn('````text', text)
        self.assertNotIn('<script>attack', (out/'report.html').read_text(encoding='utf-8'))

    def test_missing_cache_does_not_become_none_record(self):
        from openpyxl import Workbook
        book=Workbook(); book.active['A1']='=1+2'; book.save(self.root/'book.xlsx')
        evidence=sw.extract(self.root)
        self.assertFalse(any(r['text']=='None' for r in evidence['records']))
        self.assertTrue(any('cached result missing' in str(f['notices']) for f in evidence['files']))

    def test_array_formula_snapshot_is_stable(self):
        from openpyxl import Workbook
        from openpyxl.worksheet.formula import ArrayFormula
        book=Workbook(); book.active['A1']=ArrayFormula(ref='A1:A2', text='=SUM(B1:B2)'); book.save(self.root/'array.xlsx')
        a,b=sw.extract(self.root),sw.extract(self.root)
        self.assertEqual(a,b)
        self.assertTrue(any('SUM(B1:B2)' in r['text'] for r in a['records']))
        self.assertFalse(any('object at 0x' in r['text'] for r in a['records']))

    def test_empty_file_explains_coverage(self):
        (self.root/'empty.txt').write_text('')
        item=next(f for f in sw.extract(self.root)['files'] if f['path']=='empty.txt')
        self.assertEqual(item['status'],'partial'); self.assertTrue(item['notices'])

    def test_cache_directory_pruned_before_descent(self):
        folder=self.root/'node_modules'; folder.mkdir(); (folder/'private.txt').write_text('do not read')
        ev=sw.extract(self.root)
        self.assertFalse(any('do not read' in r['text'] for r in ev['records']))
        self.assertEqual(next(f['status'] for f in ev['files'] if f['path']=='node_modules'),'excluded')

    def test_inventory_limit_visible(self):
        for n in range(5): (self.root/f'{n}.txt').write_text('budget')
        with patch.object(sw,'MAX_ENTRIES',2): ev=sw.extract(self.root)
        self.assertTrue(any('Inventory' in str(f['notices']) or 'inventory' in str(f['notices']) for f in ev['files']))

    def test_text_limit_stops_extraction(self):
        with patch.object(sw,'MAX_TEXT',10): ev=sw.extract(self.root)
        self.assertEqual(ev['records'],[])
        self.assertEqual(ev['files'][0]['status'],'partial')

    def test_sparse_xlsx_dimensions_bounded(self):
        from openpyxl import Workbook
        book=Workbook(); book.active['XFD1048576']='far away'; book.save(self.root/'huge.xlsx')
        ev=sw.extract(self.root)
        item=next(f for f in ev['files'] if f['path']=='huge.xlsx')
        self.assertEqual(item['status'],'partial')
        self.assertIn('dimensions',str(item['notices']))

    def test_link_source_and_report_rejected(self):
        link=Path(self.tmp.name)/'linked.md'
        try: link.symlink_to(self.root/'note.md')
        except OSError: self.skipTest('Platform does not permit symlink creation')
        with self.assertRaises(ValueError): sw.extract(link)
        out=Path(self.tmp.name)/'out'; out.mkdir(); (out/'report.md').symlink_to(self.root/'note.md')
        with self.assertRaises(ValueError): sw.render(self.ev,self.findings,out,overwrite=True)

    def test_bad_shape_rejected(self):
        self.findings['findings']='not an array'
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)

    def test_review_correction_requires_text(self):
        self.findings['findings'][0].update(classification='needs_review',correction={'not':'text'})
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)

    def test_duplicate_findings_rejected(self):
        self.findings['findings'] *= 2
        with self.assertRaises(ValueError): sw.validate(self.ev,self.findings)

    def test_invalid_decimal_values(self):
        for value in ['NaN','Infinity','1e9999','1_000',True,0.2]:
            with self.subTest(value=value), self.assertRaises(ValueError): sw.decimal_value(value)

    def test_repeating_division_not_rounded_into_proof(self):
        f=self.findings['findings'][0]
        f.update(classification='derived',calculation=dict(op='divide',operands=['1','3'],result='0.3333333333333333333333333333',relationship_record_ids=['r2']))
        with self.assertRaisesRegex(ValueError,'arithmetic'): sw.validate(self.ev,self.findings)

    def test_docx_unsupported_content_always_disclosed(self):
        from docx import Document
        doc=Document(); doc.add_paragraph('Body'); doc.save(self.root/'basic.docx')
        item=next(f for f in sw.extract(self.root)['files'] if f['path']=='basic.docx')
        self.assertEqual(item['status'],'partial'); self.assertTrue(item['notices'])

    def test_cli_dotdot_cannot_write_inside_source(self):
        out=self.root/'new-output'
        result=subprocess.run([sys.executable,sw.__file__,'extract',str(self.root/'..'/'source'),'--out',str(out)],capture_output=True,text=True)
        self.assertNotEqual(result.returncode,0)
        self.assertIn('outside source',result.stderr)
        self.assertFalse(out.exists())

    def test_dimensionless_xlsx_has_iteration_limit(self):
        from openpyxl import Workbook
        book=Workbook()
        for n in range(1,9): book.active.cell(n,1,n)
        path=self.root/'nodim.xlsx'; book.save(path)
        buffer=io.BytesIO()
        with zipfile.ZipFile(path) as original, zipfile.ZipFile(buffer,'w') as rewritten:
            for item in original.infolist():
                data=original.read(item.filename)
                if item.filename.startswith('xl/worksheets/'):
                    data=re.sub(rb'<dimension[^>]*/>',b'',data)
                rewritten.writestr(item,data)
        path.write_bytes(buffer.getvalue())
        with patch.object(sw,'MAX_CELLS',3): evidence=sw.extract(self.root)
        item=next(f for f in evidence['files'] if f['path']=='nodim.xlsx')
        self.assertEqual(item['status'],'partial')
        self.assertIn('actual iteration',str(item['notices']))
